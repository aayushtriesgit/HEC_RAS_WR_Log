"""Write a HEC-RAS project inventory to an Excel workbook."""

from __future__ import annotations

from pathlib import Path
from typing import List, Optional, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import ProjectInventory
from .parsers import file_titles

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="44546A")

PROJECT_HEADERS = ["Field", "Value"]
FILES_HEADERS = ["File Name", "Extension", "File Type", "Title", "Full Path"]
PLANS_HEADERS = [
    "Plan File",
    "Plan Title",
    "Short ID",
    "Geometry File",
    "Geometry Title",
    "Flow File",
    "Flow Title",
    "Simulation Date",
    "Computation Interval",
    "Output Interval",
    "Program Version",
]
GEOM_HEADERS = [
    "Geometry File",
    "Geometry Title",
    "Associated Terrain",
    "Program Version",
    "Full Path",
]
TERRAIN_HEADERS = ["Terrain Name", "Terrain File", "Priority"]


def _write_table(
    sheet: Worksheet,
    headers: Sequence[str],
    rows: List[Sequence[object]],
    style_header: bool = True,
) -> None:
    sheet.append(list(headers))
    if style_header:
        for cell in sheet[sheet.max_row]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
    for row in rows:
        sheet.append(list(row))
    for col_idx, header in enumerate(headers, start=1):
        width = max(
            [len(str(header))]
            + [len(str(row[col_idx - 1] or "")) for row in rows]
        )
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(width + 3, 60)


def build_rows(inventory: ProjectInventory) -> dict:
    """Build the row data for every sheet, keyed by sheet name."""
    titles = file_titles(inventory)

    project_rows = [
        ("Project Title", inventory.title),
        ("Project File", inventory.project_file.name if inventory.project_file else ""),
        ("Project Folder", str(inventory.folder)),
        ("Units", inventory.units),
        ("Current Plan", inventory.current_plan),
        ("Projection", inventory.projection),
        ("Description", inventory.description),
        ("Number of Plans", len(inventory.plans)),
        ("Number of Geometries", len(inventory.geometries)),
        ("Number of Flow Files", len(inventory.flows)),
        ("Number of Terrains", len(inventory.terrains)),
    ]

    file_rows = [
        (f.name, f.extension, f.file_type, f.title, str(f.path))
        for f in inventory.all_files()
    ]

    plan_rows = []
    for plan in inventory.plans:
        plan_rows.append(
            (
                plan.file.name,
                plan.title,
                plan.short_id,
                plan.geom_ext,
                titles.get(plan.geom_ext.lower(), ""),
                plan.flow_ext,
                titles.get(plan.flow_ext.lower(), ""),
                plan.simulation_date,
                plan.computation_interval,
                plan.output_interval,
                plan.program_version,
            )
        )

    geom_rows = [
        (
            g.file.name,
            g.title,
            g.terrain_name,
            g.program_version,
            str(g.file.path),
        )
        for g in inventory.geometries
    ]

    terrain_rows = [(t.name, t.filename, t.priority) for t in inventory.terrains]

    return {
        "Project": (PROJECT_HEADERS, project_rows),
        "Files": (FILES_HEADERS, file_rows),
        "Plans": (PLANS_HEADERS, plan_rows),
        "Geometries": (GEOM_HEADERS, geom_rows),
        "Terrains": (TERRAIN_HEADERS, terrain_rows),
    }


def write_inventory(
    inventory: ProjectInventory,
    output: Path,
    template: Optional[Path] = None,
) -> Path:
    """Write the inventory to ``output``.

    If ``template`` is given, the workbook is loaded from it and data rows are
    appended to any sheet whose name matches an inventory section (Project,
    Files, Plans, Geometries, Terrains); headers are assumed to already exist
    in the template. Missing sheets are created with default headers.
    """
    sections = build_rows(inventory)

    if template is not None:
        workbook = load_workbook(template)
        for sheet_name, (headers, rows) in sections.items():
            if sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                for row in rows:
                    sheet.append(list(row))
            else:
                sheet = workbook.create_sheet(sheet_name)
                _write_table(sheet, headers, rows)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)
        for sheet_name, (headers, rows) in sections.items():
            sheet = workbook.create_sheet(sheet_name)
            _write_table(sheet, headers, rows)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output
